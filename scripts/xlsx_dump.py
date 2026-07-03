#!/usr/bin/env python3
"""Dump non-empty cells of an xlsx to stdout, sheet by sheet.
Usage: python3 scripts/xlsx_dump.py <xlsx_path> [max_rows_per_sheet]
"""
import sys
from openpyxl import load_workbook

def main():
    path = sys.argv[1]
    maxr = int(sys.argv[2]) if len(sys.argv) > 2 else 400
    wb = load_workbook(path, data_only=True, read_only=True)
    print(f"### {path}")
    print(f"sheets: {wb.sheetnames}")
    for ws in wb.worksheets:
        print(f"\n===== SHEET: {ws.title} =====")
        rn = 0
        for row in ws.iter_rows():
            nonempty = [c for c in row if c.value not in (None, "")]
            if not nonempty:
                continue
            rn += 1
            if rn > maxr:
                print("  ...[truncated]")
                break
            rownum = nonempty[0].row
            line = "  ".join(f"{c.coordinate}={c.value!r}" for c in nonempty)
            print(f"[{rownum}] {line}")

if __name__ == "__main__":
    main()
